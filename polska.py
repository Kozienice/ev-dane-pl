"""
polska.py (dawniej pobierz.py) - warstwa pobierania danych dla zadania "EV Polska - przeglad tygodniowy".

WERSJA 0.2 (2026-09-29)
- PZPM eRejestracje: 2 ostatnie miesiace (XLSX + komunikat PDF), zrzut arkuszy do CSV,
  parser rankingow BEV (marki, modele; narastajaco) i arkusza "Ogolem" (miesiac + YTD wg napedu),
  kontrola spojnosci (suma rankingu = Elektryczne YTD w "Ogolem"),
  wyliczenie miesiecznych rejestracji per marka = YTD(m) - YTD(m-1), tylko z plikow PZPM.
- CEPiK: jedna proba tygodniowo (standardowe TLS) - wykrywa, czy serwer poprawil klucz DH.
- ORLEN Charge: surowy HTML cennika + wyciag linii z cenami.
- CPCA usuniete (403 z GitHuba, 29.09.2026).

Zasady:
- zadnych liczb wpisanych recznie; wszystko w data/ pochodzi z odpowiedzi zrodla albo z arytmetyki na nich
- kazdy pobrany plik ma w data/status.json: adres, czas UTC, kod HTTP, rozmiar, sha256
- blad jednego zrodla nie przerywa pozostalych; skrypt zawsze konczy sie zapisem status.json
- polaczenia wylacznie ze standardowymi ustawieniami bezpieczenstwa TLS
"""

import csv
import hashlib
import json
import re
import sys
import traceback
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urljoin

import requests

DATA = Path("data")
RAW = DATA / "raw"
CSV = DATA / "csv"
TAB = DATA / "tabele"
UA = "ev-dane-pl/0.2-polska (+https://github.com/Kozienice/ev-dane-pl)"
TIMEOUT = 60

status = {
    "uruchomienie_utc": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
    "wersja_skryptu": "0.2",
    "zrodla": {},
}


# ---------------------------------------------------------------- narzedzia

def sesja():
    s = requests.Session()
    s.headers["User-Agent"] = UA
    return s


def zapisz(sciezka: Path, tresc: bytes, zrodlo: str, url: str, kod: int):
    sciezka.parent.mkdir(parents=True, exist_ok=True)
    sciezka.write_bytes(tresc)
    status["zrodla"].setdefault(zrodlo, {}).setdefault("pliki", []).append({
        "plik": str(sciezka),
        "url": url,
        "http": kod,
        "bajty": len(tresc),
        "sha256": hashlib.sha256(tresc).hexdigest(),
        "pobrano_utc": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
    })


def notuj(zrodlo: str, klucz: str, wartosc):
    status["zrodla"].setdefault(zrodlo, {})[klucz] = wartosc


def get(s, url, zrodlo):
    try:
        r = s.get(url, timeout=TIMEOUT)
        notuj(zrodlo, f"GET {url}", {"http": r.status_code, "bajty": len(r.content),
                                     "content_type": r.headers.get("Content-Type")})
        return r
    except Exception as e:
        notuj(zrodlo, f"GET {url}", {"blad": f"{type(e).__name__}: {e}"[:500]})
        return None


def zapisz_csv(sciezka: Path, naglowek, wiersze):
    sciezka.parent.mkdir(parents=True, exist_ok=True)
    with sciezka.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(naglowek)
        w.writerows(wiersze)


def liczba(v):
    """Komorka -> float albo None. Nie zgaduje: tekst nieliczbowy = None."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace(" ", "").replace(" ", "").replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


# ---------------------------------------------------------------- PZPM: pobieranie

MIESIACE = ["STYCZEN", "LUTY", "MARZEC", "KWIECIEN", "MAJ", "CZERWIEC",
            "LIPIEC", "SIERPIEN", "WRZESIEN", "PAZDZIERNIK", "LISTOPAD", "GRUDZIEN"]
BAZA_PZPM = "https://www1.pzpm.org.pl"
INDEKS_PZPM = BAZA_PZPM + "/pl/Elektromobilnosc/eRejestracje"


def strony_miesieczne(html: str, rok: int):
    """Tylko linki postaci /pl/Elektromobilnosc/eRejestracje/<MIESIAC>-<ROK><cyfry?> (slugi PZPM bywaja z literowka, np. STYCZEN-20262).
    Pomija przelaczniki jezyka, wersje do druku i podstrony plikow."""
    wynik = {}
    for href in set(re.findall(r'href="(/pl/Elektromobilnosc/eRejestracje/[^"/]+)"', html)):
        slug = href.rsplit("/", 1)[-1].upper()
        for i, m in enumerate(MIESIACE, start=1):
            if re.fullmatch(rf"{m}-{rok}\d?", slug):
                wynik[i] = urljoin(BAZA_PZPM, href)
    return dict(sorted(wynik.items()))


def pzpm():
    z = "pzpm"
    s = sesja()
    r = get(s, INDEKS_PZPM, z)
    if r is None or r.status_code != 200:
        return
    zapisz(RAW / "pzpm" / "indeks_eRejestracje.html", r.content, z, INDEKS_PZPM, r.status_code)
    rok = datetime.now(timezone.utc).year
    mies = strony_miesieczne(r.text, rok)
    notuj(z, "strony_miesieczne", {f"{rok}-{k:02d}": v for k, v in mies.items()})

    pobrane = []
    for nr, url in list(mies.items())[-2:]:
        okres = f"{rok}-{nr:02d}"
        rs = get(s, url, z)
        if rs is None or rs.status_code != 200:
            continue
        xlsx = sorted({urljoin(url, h) for h in re.findall(r'href="([^"]+\.xlsx)"', rs.text)})
        pdf = sorted({urljoin(url, h) for h in re.findall(r'href="([^"]+\.pdf)"', rs.text)})
        notuj(z, f"pliki_{okres}", {"xlsx": xlsx, "pdf": pdf})
        for u in xlsx[:1]:
            rx = get(s, u, z)
            if rx is not None and rx.status_code == 200:
                p = RAW / "pzpm" / f"{okres}_tabele.xlsx"
                zapisz(p, rx.content, z, u, rx.status_code)
                xlsx_do_csv(p, CSV / "pzpm" / okres, z)
                pobrane.append((okres, p))
        for u in pdf:
            if "informacja" in u.lower():
                rp = get(s, u, z)
                if rp is not None and rp.status_code == 200:
                    zapisz(RAW / "pzpm" / f"{okres}_informacja.pdf", rp.content, z, u, rp.status_code)

    wyniki = {}
    for okres, p in pobrane:
        try:
            wyniki[okres] = parsuj_pzpm(p, okres, z)
        except Exception:
            notuj(z, f"blad_parsera_{okres}", traceback.format_exc()[-1500:])
    if len(wyniki) == 2:
        miesiac_z_roznicy(wyniki, z)


def xlsx_do_csv(sciezka: Path, katalog: Path, z: str):
    """Zrzut KAZDEGO arkusza do CSV 1:1 (bez interpretacji)."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(sciezka, data_only=True, read_only=True)
        arkusze = []
        for ws in wb.worksheets:
            nazwa = re.sub(r"[^\w\-]+", "_", ws.title)[:60]
            wiersze = [["" if v is None else v for v in row] for row in ws.iter_rows(values_only=True)]
            out = katalog / f"{nazwa}.csv"
            out.parent.mkdir(parents=True, exist_ok=True)
            with out.open("w", newline="", encoding="utf-8") as f:
                csv.writer(f).writerows(wiersze)
            arkusze.append({"arkusz": ws.title, "wiersze": len(wiersze), "plik": str(out)})
        notuj(z, f"arkusze_{sciezka.stem}", arkusze)
    except Exception as e:
        notuj(z, f"blad_xlsx_{sciezka.stem}", f"{type(e).__name__}: {e}"[:500])


# ---------------------------------------------------------------- PZPM: parsery

def wiersze_arkusza(p: Path, nazwa: str):
    from openpyxl import load_workbook
    wb = load_workbook(p, data_only=True, read_only=True)
    if nazwa not in wb.sheetnames:
        return None
    return [list(r) for r in wb[nazwa].iter_rows(values_only=True)]


def ranking_bev(rows):
    """Blok 'Elektryczne: BEV' w arkuszu 'Osobowe - rankingi'.
    Uklad (sprawdzony na pliku 08.2026): kol.1 pozycja, 2 marka, 3 YTD biezacy, 4 udzial, 5 YTD poprzedni, 6 udzial, 7 zmiana r/r;
    kol.9 pozycja, 10 model, 11-15 analogicznie. Blok konczy wiersz 'Razem'.
    Uklad jest WERYFIKOWANY naglowkami - gdy sie zmieni, parser zglasza blad zamiast zgadywac."""
    start = next((i for i, r in enumerate(rows) if any(isinstance(c, str) and c.strip() == "Elektryczne: BEV" for c in r)), None)
    if start is None:
        raise ValueError("brak bloku 'Elektryczne: BEV'")
    hdr = next(i for i in range(start, start + 8) if rows[i][2] == "Marka" and rows[i][10] == "Model")
    okres_opis = str(rows[hdr][3]).strip()
    lata = (str(rows[hdr + 1][3]).strip(), str(rows[hdr + 1][5]).strip())
    marki, modele, razem = [], [], None
    for r in rows[hdr + 3:]:
        etyk = str(r[1]).strip() if r[1] is not None else ""
        if etyk == "Razem":
            razem = (liczba(r[3]), liczba(r[5]), liczba(r[7]))
            break
        if etyk.startswith("Razem 1-") or etyk == "Pozostałe":
            marki.append([etyk, None, liczba(r[3]), liczba(r[4]), liczba(r[5]), liczba(r[6]), liczba(r[7])])
            modele.append([str(r[9]).strip(), None, liczba(r[11]), liczba(r[12]), liczba(r[13]), liczba(r[14]), liczba(r[15])])
            continue
        if liczba(r[1]) is not None:
            marki.append([str(r[2]).strip(), int(liczba(r[1])), liczba(r[3]), liczba(r[4]), liczba(r[5]), liczba(r[6]), liczba(r[7])])
            modele.append([str(r[10]).strip(), int(liczba(r[9])), liczba(r[11]), liczba(r[12]), liczba(r[13]), liczba(r[14]), liczba(r[15])])
    if razem is None:
        raise ValueError("brak wiersza 'Razem' w bloku BEV")
    return okres_opis, lata, marki, modele, razem


def ogolem(rows):
    """Arkusz 'Ogolem': segment OSOBOWE, wiersze wg napedu. Kol.2 miesiac, 3 udzial, 4 r/r, 5 YTD, 6 udzial YTD, 7 r/r YTD."""
    naglowek = next(r for r in rows if any(isinstance(c, str) and "Zmiana" in c for c in r))
    mies_opis, ytd_opis = str(naglowek[2]).strip(), str(naglowek[5]).strip()
    wyn, w_osobowych = [], False
    for r in rows:
        etyk = str(r[1]).strip() if r[1] is not None else ""
        if etyk == "OSOBOWE":
            w_osobowych = True
        elif etyk.isupper() and len(etyk) > 3 and etyk != "OSOBOWE":
            if w_osobowych:
                break
        if w_osobowych and etyk and etyk != "w tym:":
            wyn.append([etyk, liczba(r[2]), liczba(r[3]), liczba(r[4]), liczba(r[5]), liczba(r[6]), liczba(r[7])])
    return mies_opis, ytd_opis, wyn


def parsuj_pzpm(p: Path, okres: str, z: str):
    rows = wiersze_arkusza(p, "Osobowe - rankingi")
    if rows is None:
        raise ValueError("brak arkusza 'Osobowe - rankingi'")
    opis, lata, marki, modele, razem = ranking_bev(rows)
    nag = ["pozycja_lub_etykieta", "pozycja", f"ytd_{lata[0]}", f"udzial_{lata[0]}", f"ytd_{lata[1]}", f"udzial_{lata[1]}", "zmiana_rr"]
    nag[0] = "marka"
    zapisz_csv(TAB / f"pzpm_bev_marki_ytd_{okres}.csv", nag, marki)
    nag[0] = "model"
    zapisz_csv(TAB / f"pzpm_bev_modele_ytd_{okres}.csv", nag, modele)

    og = wiersze_arkusza(p, "Ogółem")
    kontrola = {"okres_rankingu": opis, "razem_bev_ranking": razem[0]}
    if og is not None:
        mies_opis, ytd_opis, wiersze = ogolem(og)
        zapisz_csv(TAB / f"pzpm_osobowe_naped_{okres}.csv",
                   ["naped", "miesiac", "udzial_miesiac", "rr_miesiac", "ytd", "udzial_ytd", "rr_ytd"], wiersze)
        bev = next((w for w in wiersze if w[0] == "Elektryczne"), None)
        kontrola.update({"miesiac_opis": mies_opis, "ytd_opis": ytd_opis,
                         "elektryczne_ytd_ogolem": bev[4] if bev else None,
                         "elektryczne_miesiac_ogolem": bev[1] if bev else None})
        kontrola["zgodnosc_sum"] = (bev is not None and razem[0] == bev[4])
    notuj(z, f"kontrola_{okres}", kontrola)
    return {"marki": {m[0]: m[2] for m in marki if m[1] is not None}, "razem": razem[0]}


def miesiac_z_roznicy(wyniki, z):
    """Miesieczne rejestracje per marka = YTD(m) - YTD(m-1), oba z plikow PZPM.
    Tylko marki obecne w TOP obu miesiecy; korekty CEP miedzy publikacjami moga zaburzyc wynik -> flaga w pliku."""
    (o1, w1), (o2, w2) = sorted(wyniki.items())
    wiersze = []
    for marka, ytd2 in w2["marki"].items():
        ytd1 = w1["marki"].get(marka)
        if ytd1 is None or ytd2 is None:
            wiersze.append([marka, ytd2, ytd1, None, "brak w TOP poprzedniego miesiaca"])
        else:
            d = ytd2 - ytd1
            wiersze.append([marka, ytd2, ytd1, d, f"{int(ytd2)}-{int(ytd1)}" + ("; UJEMNE -> korekta CEP" if d < 0 else "")])
    razem_d = (w2["razem"] - w1["razem"]) if w1["razem"] is not None and w2["razem"] is not None else None
    wiersze.append(["RAZEM BEV", w2["razem"], w1["razem"], razem_d, "kontrola: porownaj z 'Elektryczne' miesiac w pzpm_osobowe_naped"])
    zapisz_csv(TAB / f"pzpm_bev_marki_miesiac_{o2}.csv",
               ["marka", f"ytd_{o2}", f"ytd_{o1}", "miesiac_wyliczony", "formula_uwagi"], wiersze)
    notuj(z, "miesiac_z_roznicy", {"okres": o2, "baza": o1, "razem_wyliczony": razem_d})


# ---------------------------------------------------------------- CEPiK (sonda tygodniowa)

def cepik():
    z = "cepik"
    url = "https://api.cepik.gov.pl/slowniki"
    r = get(sesja(), url, z)
    if r is not None:
        zapisz(RAW / "cepik" / "slowniki.json", r.content[:2_000_000], z, url, r.status_code)


# ---------------------------------------------------------------- ORLEN Charge (cennik)

def orlen():
    z = "orlen_charge"
    url = "https://orlencharge.pl/cennik"
    r = get(sesja(), url, z)
    if r is None or r.status_code != 200:
        return
    zapisz(RAW / "orlen" / "cennik.html", r.content, z, url, r.status_code)
    tekst = re.sub(r"<[^>]+>", " ", r.text)
    tekst = re.sub(r"\s+", " ", tekst)
    fragmenty = sorted(set(m.group(0).strip() for m in re.finditer(r".{0,80}zł\s*/\s*(kWh|min).{0,40}", tekst)))
    zapisz_csv(TAB / "orlen_cennik_fragmenty.csv", ["fragment_tekstu"], [[f] for f in fragmenty])
    notuj(z, "fragmenty_z_cenami", len(fragmenty))


# ---------------------------------------------------------------- main

def main():
    for nazwa, f in (("pzpm", pzpm), ("cepik", cepik), ("orlen_charge", orlen)):
        try:
            f()
        except Exception:
            notuj(nazwa, "wyjatek", traceback.format_exc()[-1500:])
    DATA.mkdir(exist_ok=True)
    (DATA / "status.json").write_text(json.dumps(status, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(status, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    sys.exit(main())
